"""One actual Office job, then read-only restoration witnesses inside Server.

Run through the public FILE adapter. No retry of POST, no provider invocation,
and no reconstruction of saved options when checking an existing job.
"""
import hashlib
import http.client
import io
import json
import os
import re
import ssl
import sys
import time
import uuid

from docx import Document

SENTINEL = 'LT197 owned cold state DOCX sentinel'
PHASE = 'configuration'


def retained_options(status):
    metadata = json.loads(status['metadata_json'])
    if (not isinstance(metadata, dict) or set(metadata) != {'allow_semantic_rebuild', 'acceptance'} or
            metadata['allow_semantic_rebuild'] is not False or not isinstance(metadata['acceptance'], dict) or
            type(metadata['acceptance'].get('passed')) is not bool):
        raise ValueError('explicit original admission options changed')
    # Keep the complete enriched metadata in statusSnapshot for byte-equivalent
    # restoration checks; extract only the original, explicitly disabled option.
    return {'allow_semantic_rebuild': metadata['allow_semantic_rebuild']}


def _artifact(raw):
    document = Document(io.BytesIO(raw))
    if SENTINEL not in '\n'.join(paragraph.text for paragraph in document.paragraphs):
        raise ValueError('genuine DOCX sentinel unavailable')
    return hashlib.sha256(raw).hexdigest()


def verify_existing(status, raw, expected):
    """Require original terminal metadata and byte-identical parseable output."""
    if status != expected['statusSnapshot'] or status['status'] != 'done' or status['id'] != expected['jobID']:
        raise ValueError('persisted job metadata changed')
    if hashlib.sha256(raw).hexdigest() != expected['outputSHA256']:
        raise ValueError('original artifact hash changed')
    _artifact(raw)
    return {'jobID': status['id'], 'inputSHA256': expected['inputSHA256'],
            'outputSHA256': expected['outputSHA256'], 'persistedOptionsUnchanged': True,
            'genuineDOCXParsed': True, 'postRepeated': False}


def main():
    global PHASE
    mode = sys.argv[1]
    if mode not in ('create', 'verify-existing') or os.environ.get('APP_ENV') != 'production':
        raise ValueError('production create/verify-existing mode required')
    key = os.environ['APP_API_KEY']
    context = ssl.create_default_context(cafile='/run/tls/ca.pem')

    def request(method, path, body=None, content_type=None):
        connection = http.client.HTTPSConnection('remedy-server', 8000, context=context, timeout=120)
        try:
            headers = {'X-API-Key': key}
            if content_type:
                headers['Content-Type'] = content_type
            connection.request(method, path, body=body, headers=headers)
            response = connection.getresponse()
            raw = response.read(8 * 1024**2 + 1)
            if len(raw) > 8 * 1024**2:
                raise ValueError('bounded job response exceeded')
            return response.status, raw
        finally:
            connection.close()

    if mode == 'create':
        document = Document()
        document.add_heading('Synthetic retained Office job', level=1)
        document.add_paragraph(SENTINEL)
        stream = io.BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        boundary = 'lt197-' + uuid.uuid4().hex
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="allow_semantic_rebuild"\r\n\r\nfalse\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="synthetic-retained.docx"\r\n'
                'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n').encode()
        body += raw + f'\r\n--{boundary}--\r\n'.encode()
        # Exactly one submission. An ambiguous response is a failed witness,
        # never permission to resubmit or reconstruct its admission options.
        PHASE = 'single-submission'
        code, response = request('POST', '/v1/remediate', body, 'multipart/form-data; boundary=' + boundary)
        if code != 202:
            raise ValueError('Office job was not accepted')
        job_id = json.loads(response)['id']
        expected = {'jobID': job_id, 'inputSHA256': hashlib.sha256(raw).hexdigest()}
    else:
        expected = json.load(sys.stdin)
        job_id = expected['jobID']
    if not re.fullmatch(r'[a-f0-9]{32}', job_id):
        raise ValueError('exact job identity required')
    deadline = time.monotonic() + 120
    PHASE = 'terminal-status'
    while True:
        code, raw_status = request('GET', '/v1/jobs/' + job_id)
        if code != 200:
            raise ValueError('retained job unavailable')
        status = json.loads(raw_status)
        if status['status'] in ('done', 'failed'):
            break
        if mode != 'create' or time.monotonic() >= deadline:
            raise ValueError('terminal job deadline or restoration state failed')
        time.sleep(.2)
    if status['status'] != 'done' or status['kind'] != 'remediate_office':
        raise ValueError('actual Office remediation did not complete')
    PHASE = 'original-download'
    code, result = request('GET', '/v1/jobs/' + job_id + '/result')
    if code != 200:
        raise ValueError('retained Office download unavailable')
    if mode == 'create':
        PHASE = 'retained-options'
        options = retained_options(status)
        PHASE = 'genuine-artifact'
        expected.update(outputSHA256=_artifact(result), statusSnapshot=status)
        # The complete synthetic terminal snapshot is PRIVATE input for the
        # caller's subsequent restoration checks, never an uploaded raw log.
        print(json.dumps({**expected, 'options': options, 'postCount': 1, 'genuineDOCXParsed': True}))
    else:
        print(json.dumps(verify_existing(status, result, expected)))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        # Provider bodies, API keys and file contents never become diagnostics.
        print(json.dumps({'actualOfficeWitnessFailed': True, 'phase': PHASE, 'errorType': type(error).__name__,
                          'postRetried': False}), file=sys.stderr)
        raise SystemExit('actual Office job witness failed; no automatic submission retry') from None

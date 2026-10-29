"""Inspect calibration evidence without executing any judge or provider.

Schema labels and matching hashes cannot establish readable artifacts or human
review. An operator-controlled approval is a separate trust input; this module
never creates one or promotes synthetic evidence.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from typing import Any


@dataclass
class CorpusEvidence:
    errors: list[str] = field(default_factory=list)
    digest: str = ""
    readable_sources: int = 0
    synthetic_records: int = 0
    approved: bool = False


def artifact_readability_errors(path: Path, fmt: str) -> list[str]:
    """Actually parse a PDF or modern Office artifact; no accessibility claim."""
    try:
        if fmt == "pdf":
            import pikepdf
            with pikepdf.open(path) as document:
                if not document.pages:
                    raise ValueError("PDF has no pages")
        elif fmt == "docx":
            from docx import Document
            Document(path)
        elif fmt == "pptx":
            from pptx import Presentation
            Presentation(path)
        elif fmt == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            workbook.close()
        else:
            return ["unsupported artifact format"]
    except Exception as exc:
        return [f"artifact is not parser-readable ({type(exc).__name__})"]
    return []


def _artifact_path(value: str, root: Path) -> Path:
    from tools.annotate_corpus import _resolve_artifact_path
    resolved = _resolve_artifact_path(value, root=root)
    if resolved is None or not resolved.is_file():
        raise ValueError("artifact is missing")
    return resolved


def inspect_corpus_evidence(
    records: list[dict[str, Any]], *, root: Path, judge_results: Path | None = None,
    approval_path: Path | None = None, judge_comparisons: Path | None = None,
) -> CorpusEvidence:
    evidence = CorpusEvidence()
    digest_items = []
    if not records:
        evidence.errors.append("no corpus evidence")
    for record in sorted(records, key=lambda item: (str(item.get("format", "")), str(item.get("doc_id", "")))):
        identifier = str(record.get("doc_id", "unknown"))
        fmt = str(record.get("format", ""))
        provenance = record.get("provenance") if isinstance(record.get("provenance"), dict) else {}
        attribution = str(record.get("annotator", "")).strip()
        text = " ".join([attribution, str(provenance.get("candidate_seed_notes", "")),
                         str(provenance.get("gold_standard_source", ""))])
        synthetic = bool(re.search(r"synthetic|bootstrap|placeholder|automated|unreviewed", text, re.I))
        if synthetic:
            evidence.synthetic_records += 1
            evidence.errors.append(f"{identifier}: synthetic/bootstrap or unreviewed evidence cannot establish human calibration")
        if not attribution or provenance.get("human_verified") is not True or provenance.get("gold_standard_source") != "human_specialist":
            evidence.errors.append(f"{identifier}: attributable specialist review is missing")
        hashes = record.get("artifact_hashes") or {}
        artifacts = [("source", str(record.get("source_path", "")), hashes.get("source_sha256")),
                     ("gold", str(record.get("gold_remediation_path", "")), hashes.get("gold_remediation_sha256"))]
        artifacts += [("known_bad", str(value), (hashes.get("known_bad_sha256") or {}).get(value))
                      for value in record.get("known_bad_artifact_paths", [])]
        if not record.get("known_bad_artifact_paths"):
            evidence.errors.append(f"{identifier}: known-bad evidence is missing")
        for role, value, expected in artifacts:
            try:
                path = _artifact_path(value, root)
                actual = hashlib.sha256(path.read_bytes()).hexdigest()
                if actual != expected:
                    evidence.errors.append(f"{identifier}/{role}: artifact hash mismatch or missing hash")
                errors = artifact_readability_errors(path, fmt)
                evidence.errors.extend(f"{identifier}/{role}: {message}" for message in errors)
                if role == "source" and not errors:
                    evidence.readable_sources += 1
                digest_items.append((identifier, role, value, actual))
            except (OSError, ValueError):
                evidence.errors.append(f"{identifier}/{role}: artifact is missing or unreadable")
        digest_items.append((identifier, "annotation", json.dumps(record, sort_keys=True, separators=(",", ":"))))
    if judge_results is not None:
        try:
            digest_items.append(("judge_results", hashlib.sha256(judge_results.read_bytes()).hexdigest()))
        except OSError:
            evidence.errors.append("attributable judge-result evidence is missing")
    else:
        evidence.errors.append("persisted attributable judge-result evidence is required")
    if judge_comparisons is not None:
        try:
            digest_items.append(("judge_comparisons", hashlib.sha256(judge_comparisons.read_bytes()).hexdigest()))
        except OSError:
            evidence.errors.append("attributable judge-comparison evidence is missing")
    evidence.digest = hashlib.sha256(json.dumps(digest_items, sort_keys=True).encode()).hexdigest()
    if approval_path is None:
        evidence.errors.append("operator-controlled human-review approval is unavailable")
    else:
        try:
            approval = json.loads(approval_path.read_text())
            if not isinstance(approval, dict) or not isinstance(approval.get("reviewed_at"), str):
                raise ValueError("malformed approval")
            reviewer = approval["reviewer_id"]
            reviewed_at = datetime.fromisoformat(approval["reviewed_at"].replace("Z", "+00:00"))
            valid = (
                type(approval.get("schema_version")) is int and approval["schema_version"] == 1 and approval.get("classification") == "human_reviewed"
                and approval.get("corpus_digest") == evidence.digest
                and isinstance(reviewer, str) and bool(reviewer.strip())
                and not re.search(r"synthetic|bootstrap|placeholder|automated|unreviewed", reviewer, re.I)
                and isinstance(approval.get("evidence_reference"), str) and bool(approval["evidence_reference"].strip())
                and reviewed_at.tzinfo is not None and reviewed_at <= datetime.now(timezone.utc)
            )
            if not valid:
                raise ValueError("approval does not bind attributable review to this corpus")
        except (OSError, ValueError, TypeError, KeyError):
            evidence.errors.append("human-review approval is invalid or does not match current evidence")
    evidence.approved = not evidence.errors
    return evidence

"""Actual production HTTP/store/queue hard exits and real Office recovery.

Only persistence/publication fault points and a pre-dispatch capacity barrier
are injected. Cold recovery runs the real native Office engine in its owned
process. No provider, live service, migration or deletion is used.
"""
from contextlib import asynccontextmanager
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import signal
import shutil
import socket
import sqlite3
import subprocess
import sys
import time

import httpx
import pytest

# Direct fault subprocesses run from source in the non-installed shipping test image.
if __name__ == "__main__":
    source_root = Path(__file__).resolve().parents[2]
    sys.path[:0] = [str(source_root / "src"), str(source_root)]

from backend.app.config import Settings
from backend.app.supervisor import supported

pytestmark = pytest.mark.skipif(not supported(), reason="requires the real Linux pidfd/subreaper profile")


@pytest.fixture(autouse=True)
def retain_synthetic_http_evidence(tmp_path, request):
    yield  # all Owner finally cleanup runs before evidence I/O
    output = os.environ.get("REMEDY_SANDBOX_AUDIT_DIR")
    if output:
        destination = Path(output) / "submission-http" / request.node.name
        shutil.copytree(tmp_path, destination, dirs_exist_ok=True)
        receipt = {"scope": "Post-test synthetic diagnostic snapshot, not backup/restore certification"}
        try:
            receipt["jobs"] = rows(tmp_path) if (tmp_path / "jobs.db").exists() else []
        except Exception as error:
            receipt["diagnostic_error"] = type(error).__name__
        (destination / "receipt.json").write_text(json.dumps(receipt, indent=2))
        # pytest temp directories are private. Only this synthetic diagnostic
        # COPY is readable by the CI artifact uploader outside the test UID.
        for item in [destination, *destination.rglob("*")]:
            item.chmod(0o755 if item.is_dir() else 0o644)


def settings_for(directory):
    return Settings(app_env="production", docs_enabled=False, api_key="fixture-key",
        ollama_api_key="fixture-provider-key", cors_allow_origins=["https://fixture.invalid"],
        job_store_path=directory / "jobs.db", job_dir=directory / "jobs", backup_dir=directory / "backups",
        job_total_timeout_seconds=60, quality_require_calibration=True,
        quality_experiment_store_path=directory / "quality.db", quality_corpus_root_path=directory / "absent-corpus")


def serve(directory, point, descriptor):
    import uvicorn
    from backend.app.main import create_app
    app = create_app(settings_for(directory))
    store, worker = app.state.store, app.state.worker
    if point in {"db-before-commit", "db-after-commit"}:
        from tests.unit.submission_faults import exit_at_intent_commit
        exit_at_intent_commit(store, point)
    if point == "rename":
        rename = Path.rename
        def crash_rename(path, destination):
            result = rename(path, destination)
            if path.name.startswith("_staging-"):
                os._exit(73)
            return result
        Path.rename = crash_rename
    if point == "ready":
        advance = store._advance_upload
        async def crash_ready(*args):
            result = await advance(*args)
            os._exit(73)
        store._advance_upload = crash_ready
    published = False
    enqueue = worker.enqueue
    async def publication(*args):
        nonlocal published
        if point == "enqueue-before":
            os._exit(73)
        await enqueue(*args)
        published = True
        if point == "enqueue-after":
            os._exit(73)
    worker.enqueue = publication
    if point == "receipt":
        get = store.get
        async def crash_receipt(*args):
            result = await get(*args)
            if published:
                os._exit(73)
            return result
        store.get = crash_receipt
    begin = worker.begin_dispatch
    async def dispatch(job_id):
        await begin(job_id)
        with (directory / "dispatch.jsonl").open("a") as log:
            log.write(json.dumps({"id": job_id}) + "\n")
            log.flush()
            os.fsync(log.fileno())
        if point == "dispatch":
            os._exit(73)
    worker.begin_dispatch = dispatch
    if point not in {"recover", "dispatch"}:
        original = app.router.lifespan_context
        @asynccontextmanager
        async def held_capacity(application):
            async with original(application):
                await app.state.execution._slots.acquire()
                yield
        app.router.lifespan_context = held_capacity
    async def owned_app(scope, receive, send):
        async def observed_send(message):
            await send(message)
            if (point == "ack-sent" and scope["type"] == "http" and scope["method"] == "POST"
                    and scope["path"] == "/v1/remediate" and message["type"] == "http.response.body"
                    and not message.get("more_body", False)):
                os._exit(73)
        await app(scope, receive, observed_send)
    server = uvicorn.Server(uvicorn.Config(owned_app, lifespan="on", log_level="warning", proxy_headers=False))
    server.run(sockets=[socket.socket(fileno=descriptor)])


class Owner:
    def __init__(self, directory, point):
        listener = socket.socket()
        listener.bind(("127.0.0.1", 0))
        listener.listen(32)
        self.url = f"http://127.0.0.1:{listener.getsockname()[1]}"
        self.log = (directory / f"server-{point}-{time.monotonic_ns()}.log").open("wb")
        environment = os.environ | {
            "PYTHON_DOTENV_DISABLED": "1", "APP_ENV": "production", "APP_API_KEY": "fixture-key",
            "OLLAMA_API_KEY": "fixture-provider-key", "OLLAMA_BASE_URL": "https://fixture.invalid/v1", "CORS_ALLOW_ORIGINS": "https://fixture.invalid",
            "CXONE_INTEGRATION_ENABLED": "false", "JOB_STORE_PATH": str(directory / "import.db"),
            "JOB_DIR": str(directory / "import-jobs"), "RATE_LIMIT_UPLOADS": "10000/minute",
            "RATE_LIMIT_DEFAULT": "10000/minute",  # bounded observer polling; production defaults stay unchanged
        }
        try:
            self.process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), str(directory), point,
                str(listener.fileno())], pass_fds=(listener.fileno(),), stdout=self.log, stderr=subprocess.STDOUT,
                env=environment)
            self.handle = os.pidfd_open(self.process.pid)  # direct child remains unreaped; pin before any poll/wait
        finally:
            listener.close()
    def ready(self):
        until = time.monotonic() + 15
        while time.monotonic() < until:
            assert self.process.poll() is None, "API exited before readiness; see retained server log"
            try:
                if httpx.get(self.url + "/readyz", timeout=.5, trust_env=False).status_code == 200:
                    return
            except httpx.HTTPError:
                pass
        raise AssertionError("Actual API did not become ready within startup watchdog")
    def close(self):
        forced = False
        try:
            try:
                signal.pidfd_send_signal(self.handle, signal.SIGTERM)
            except ProcessLookupError:
                pass
            try:
                self.process.wait(timeout=15)
            except subprocess.TimeoutExpired:
                forced = True
                signal.pidfd_send_signal(self.handle, signal.SIGKILL)
                self.process.wait(timeout=5)
        finally:
            os.close(self.handle)
            self.log.close()
        return forced


def rows(directory):
    with sqlite3.connect(directory / "jobs.db") as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute("SELECT * FROM jobs")]


@pytest.mark.parametrize("point", ["db-before-commit", "db-after-commit", "rename", "ready",
    "enqueue-before", "enqueue-after", "receipt", "ack-sent", "dispatch"])
def test_real_http_hard_exit_retains_exact_submission_and_cold_recovers_safely(tmp_path, point):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Retained HTTP submission sentinel")
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    owner = Owner(tmp_path, point)
    try:
        owner.ready()
        try:
            route = "/v1/office/remediate" if point == "receipt" else "/v1/remediate"
            response = httpx.post(owner.url + route + "?quality=true", headers={"X-API-Key": "fixture-key"},
                files={"file": ("fixture.DOCX" if point == "ack-sent" else "fixture.docx", source)}, timeout=15, trust_env=False)
            assert response.status_code == 202
        except httpx.TransportError:
            pass  # a real connection was lost; never repeat the POST
        assert owner.process.wait(timeout=15) == 73
    finally:
        owner.close()
    saved = rows(tmp_path)
    assert len(saved) == (0 if point == "db-before-commit" else 1)
    output_hash = None
    for _ in range(2):
        recovered = Owner(tmp_path, "recover")
        try:
            recovered.ready()
            if not saved:
                assert rows(tmp_path) == []
                assert [path.read_bytes() for path in (tmp_path / "jobs").glob("_staging-*")] == [source]
                continue
            row = saved[0]
            assert json.loads(row["metadata_json"])["quality"] is True
            assert (tmp_path / "jobs" / row["id"] / "input.docx").read_bytes() == source
            until = time.monotonic() + 30
            while True:
                response = httpx.get(recovered.url + "/v1/jobs/" + row["id"], headers={"X-API-Key": "fixture-key"},
                                     timeout=5, trust_env=False)
                assert response.status_code == 200
                body = response.json()
                if body["status"] in {"done", "failed"}:
                    break
                assert time.monotonic() < until, body
            result = httpx.get(recovered.url + "/v1/jobs/" + row["id"] + "/result",
                               headers={"X-API-Key": "fixture-key"}, timeout=5, trust_env=False)
            if point == "dispatch":
                assert body["status"] == "failed" and body["stage"] == "interrupted"
                assert result.status_code == 409  # uncertain dispatch is not replayed
            else:
                assert body["status"] == "done" and result.status_code == 200, body
                metadata = json.loads(body["metadata_json"])
                assert metadata["quality_result_error"] == "Quality evaluation unavailable."
                assert "quality_result" not in metadata
                if point == "ready" and _ == 0:
                    capability = httpx.get(recovered.url + "/v1/quality/calibration?format=docx",
                        headers={"X-API-Key": "fixture-key"}, timeout=15, trust_env=False)
                    assert capability.status_code == 200
                    readiness = capability.json()["readiness"]
                    assert readiness["ready"] is False and readiness["capability"] == "unavailable"
                    assert readiness["evidence_errors"]
                    refused = httpx.post(recovered.url + "/v1/quality/audit/office",
                        headers={"X-API-Key": "fixture-key"}, files={"file": ("synthetic.docx", source)},
                        timeout=15, trust_env=False)
                    assert refused.status_code == 409
                    assert httpx.get(recovered.url + "/readyz", timeout=5, trust_env=False).status_code == 200
                assert "Retained HTTP submission sentinel" in [p.text for p in Document(BytesIO(result.content)).paragraphs]
                digest = hashlib.sha256(result.content).hexdigest()
                assert output_hash is None or output_hash == digest
                output_hash = digest
            after = rows(tmp_path)
            assert len(after) == 1
            observed_options = json.loads(after[0]["metadata_json"])
            for key, value in json.loads(row["metadata_json"]).items():
                assert observed_options[key] == value
            assert after[0]["deadline_at"] == row["deadline_at"]
        finally:
            assert not recovered.close(), "Cold-recovery API required forced shutdown"
    attempts = tmp_path / "dispatch.jsonl"
    assert (len(attempts.read_text().splitlines()) if attempts.exists() else 0) == (0 if not saved else 1)


@pytest.mark.skipif(os.environ.get("REMEDY_FULL_RUNTIME_TESTS") != "1", reason="requires the shipped PDF toolchain")
def test_real_http_native_pdf_keeps_download_when_calibration_unavailable(tmp_path):
    """Real API/process/PDF toolchain; no internal engine or acceptance adapter."""
    from reportlab.pdfgen.canvas import Canvas
    import fitz
    path = tmp_path / "synthetic.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 720, "Retained real PDF sentinel")
    canvas.save()
    source = path.read_bytes()
    owner = Owner(tmp_path, "recover")
    try:
        owner.ready()
        response = httpx.post(owner.url + "/v1/remediate?quality=true",
            headers={"X-API-Key": "fixture-key"}, files={"file": (path.name, source)},
            timeout=15, trust_env=False)
        assert response.status_code == 202
        job_id = response.json()["id"]
        until = time.monotonic() + 65
        while True:
            status = httpx.get(owner.url + "/v1/jobs/" + job_id,
                headers={"X-API-Key": "fixture-key"}, timeout=5, trust_env=False)
            assert status.status_code == 200
            job = status.json()
            if job["status"] in {"done", "failed"}:
                break
            assert time.monotonic() < until, job
        assert job["status"] == "done", job
        metadata = json.loads(job["metadata_json"])
        assert metadata["quality_result_error"] == "Quality calibration capability unavailable."
        result = httpx.get(owner.url + "/v1/jobs/" + job_id + "/result",
            headers={"X-API-Key": "fixture-key"}, timeout=5, trust_env=False)
        assert result.status_code == 200
        with fitz.open(stream=result.content, filetype="pdf") as document:
            assert document.page_count == 1
            assert "Retained real PDF sentinel" in document[0].get_text()
        refused = httpx.post(owner.url + "/v1/quality/audit/pdf",
            headers={"X-API-Key": "fixture-key"}, files={"file": (path.name, source)},
            timeout=15, trust_env=False)
        assert refused.status_code == 409
        assert httpx.get(owner.url + "/readyz", timeout=5, trust_env=False).status_code == 200
    finally:
        assert not owner.close(), "Native PDF API required forced shutdown"


if __name__ == "__main__":
    serve(Path(sys.argv[1]), sys.argv[2], int(sys.argv[3]))

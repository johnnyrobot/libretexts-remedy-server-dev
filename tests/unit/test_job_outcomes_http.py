"""Job contracts through the real authenticated HTTP/lifespan/worker seam.

Only external providers and controlled adapter faults are substituted. SQLite,
uploads, dispatch, terminal state and downloads are the production code.
"""

import asyncio
import json
import os
from io import BytesIO

import httpx
import pytest
import pytest_asyncio

from backend.app.supervisor import supported
pytestmark = pytest.mark.skipif(not supported(), reason="strict engine execution requires Linux amd64")

from tests.unit.office_fixtures import make_docx, make_pptx, make_xlsx


@pytest_asyncio.fixture
async def api(tmp_path, monkeypatch, request):
    for name, value in {
        "PYTHON_DOTENV_DISABLED": "1",
        "APP_ENV": "development",
        "APP_API_KEY": "fixture-key",
        "OLLAMA_API_KEY": "fixture-provider-key",
        "CXONE_INTEGRATION_ENABLED": "false",
        "JOB_STORE_PATH": str(tmp_path / "import.db"),
        "JOB_DIR": str(tmp_path / "import-jobs"),
        "JOB_BACKUP_DIR": str(tmp_path / "backups"),
        "RATE_LIMIT_DEFAULT": "10000/minute",
        "RATE_LIMIT_UPLOADS": "10000/minute",
        "OLLAMA_BASE_URL": "https://provider.invalid/v1",
        "MAX_RETRIES": "1",
    }.items():
        monkeypatch.setenv(name, value)
    from backend.app.config import Settings
    from backend.app.main import create_app

    settings = Settings(api_key="fixture-key", job_store_path=tmp_path / "jobs.db",
                        job_dir=tmp_path / "jobs", backup_dir=tmp_path / "backups",
                        quality_experiment_store_path=tmp_path / "quality.db")
    app = create_app(settings)
    # These focused adapter-fault tests need parent-local monkeypatches. They
    # retain real HTTP/store/queue orchestration but do not claim spawn coverage.
    # Unpatched artifact cases and real Chromium use production owned processes;
    # test_execution_http independently proves cancellation and capacity with
    # actual controlled spawned processes and descendants.
    adapter_fault_tests = {
        "test_unavailable_calibration_keeps_base_pdf_download",
        "test_claimed_office_success_requires_readable_artifact",
        "test_advisory_exception_keeps_output_without_disclosing_error",
        "test_render_failure_then_success_isolated_in_worker",
        "test_renderer_cannot_publish_unusable_pdf",
        "test_office_html_real_extraction_generation_after_corrupt_input",
        "test_office_html_invalid_claimed_output_fails",
        "test_queued_and_running_results_are_conflict",
        "test_generation_provider_error_then_valid_conversion",
    }
    if request.node.originalname in adapter_fault_tests:
        from backend.app.engine_service import run_job
        async def in_process_adapter_case(job):
            # This adapter owns its compute entry; JobWorker dequeue is no
            # longer dispatch. Keep the public running/queued assertions real.
            async with app.state.worker.execution._slots:
                await app.state.worker.begin_dispatch(job.id)
                await run_job(job, app.state.store, settings)
        app.state.worker.runner = in_process_adapter_case
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://fixture.invalid",
            headers={"X-API-Key": "fixture-key"},
        ) as client:
            yield client, settings


async def submit(client, route, name, data):
    response = await client.post(route, files={"file": (name, data)})
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]
    return await terminal(client, job_id)


async def terminal(client, job_id):
    async with asyncio.timeout(15):
        while True:
            response = await client.get(f"/v1/jobs/{job_id}")
            assert response.status_code == 200
            job = response.json()
            if job["status"] in {"done", "failed"}:
                return job
            await asyncio.sleep(0.01)


async def test_unavailable_calibration_keeps_base_office_download_and_refuses_quality_audit(api, tmp_path):
    client, _ = api
    path = make_docx(tmp_path / "synthetic.docx", body_paragraphs=["Retained sentinel"])
    response = await client.post("/v1/office/remediate?quality=true", files={"file": (path.name, path.read_bytes())})
    assert response.status_code == 202
    job = await terminal(client, response.json()["id"])
    assert job["status"] == "done", job
    metadata = json.loads(job["metadata_json"])
    assert metadata["quality_result_error"] == "Quality evaluation unavailable."
    assert "quality_result" not in metadata
    artifact = await client.get(f"/v1/jobs/{job['id']}/result")
    assert artifact.status_code == 200
    from docx import Document
    assert "Retained sentinel" in [paragraph.text for paragraph in Document(BytesIO(artifact.content)).paragraphs]
    quality = await client.post("/v1/quality/audit/office", files={"file": (path.name, path.read_bytes())})
    assert quality.status_code == 409
    capability = (await client.get("/v1/quality/calibration?format=docx")).json()["readiness"]
    assert capability["ready"] is False and capability["capability"] == "unavailable"
    assert capability["evidence_errors"]
    assert (await client.get("/healthz")).status_code == 200


async def test_unavailable_calibration_keeps_base_pdf_download(api, tmp_path, monkeypatch):
    """Real HTTP/store/worker and PDF download, controlled external PDF adapters."""
    from pathlib import Path
    from types import SimpleNamespace
    import shutil
    import pikepdf
    from reportlab.pdfgen.canvas import Canvas
    from backend.app import engine_service
    client, _ = api
    path = tmp_path / "synthetic.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 720, "Retained sentinel")
    canvas.save()
    def fix(source, output, **kwargs):
        shutil.copyfile(source, output)
    def acceptance(*args, **kwargs):
        return SimpleNamespace(passed=True, verapdf_result=SimpleNamespace(checked=False, passed=False))
    def report(**kwargs):
        directory = Path(kwargs["output_dir"])
        directory.mkdir(exist_ok=True)
        (directory / "report.html").write_text("<html><body>Synthetic report</body></html>")
    monkeypatch.setattr(engine_service, "fix_and_verify", fix)
    monkeypatch.setattr(engine_service, "evaluate_pdf_acceptance", acceptance)
    monkeypatch.setattr(engine_service, "generate_document_report", report)
    response = await client.post("/v1/remediate?quality=true", files={"file": (path.name, path.read_bytes())})
    assert response.status_code == 202
    job = await terminal(client, response.json()["id"])
    assert job["status"] == "done", job
    assert json.loads(job["metadata_json"])["quality_result_error"] == "Quality calibration capability unavailable."
    artifact = await client.get(f"/v1/jobs/{job['id']}/result")
    assert artifact.status_code == 200
    with pikepdf.open(BytesIO(artifact.content)) as pdf:
        assert len(pdf.pages) == 1


def office_fixture(tmp_path, suffix):
    path = tmp_path / f"valid.{suffix}"
    if suffix == "docx":
        return make_docx(path, body_paragraphs=["Retained sentinel"])
    if suffix == "pptx":
        return make_pptx(path)
    return make_xlsx(path)


@pytest.mark.parametrize("route", ["/v1/remediate", "/v1/office/remediate"])
@pytest.mark.parametrize("suffix", ["docx", "pptx", "xlsx"])
async def test_corrupt_office_fails_then_valid_job_downloads(api, tmp_path, route, suffix):
    client, _ = api
    failed = await submit(client, route, f"corrupt.{suffix}", b"PK\x03\x04broken")
    assert failed["status"] == "failed"
    assert failed["error"]
    assert "output_path" not in failed
    assert (await client.get(f"/v1/jobs/{failed['id']}/result")).status_code == 409

    path = office_fixture(tmp_path, suffix)
    valid = await submit(client, route, path.name, path.read_bytes())
    assert valid["status"] == "done", valid
    result = await client.get(f"/v1/jobs/{valid['id']}/result")
    assert result.status_code == 200
    assert "attachment" in result.headers["content-disposition"]
    if suffix == "docx":
        from docx import Document
        assert "Retained sentinel" in [p.text for p in Document(BytesIO(result.content)).paragraphs]
    elif suffix == "pptx":
        from pptx import Presentation
        assert Presentation(BytesIO(result.content)).slides[0].shapes.title.text == "Slide 1 Title"
    else:
        from openpyxl import load_workbook
        assert load_workbook(BytesIO(result.content)).active["A2"].value == "v2.1"


@pytest.mark.parametrize("route", ["/v1/remediate", "/v1/office/remediate"])
@pytest.mark.parametrize("suffix", ["doc", "PPT", "xLs"])
@pytest.mark.parametrize("quality", ["", "?quality=true"])
async def test_legacy_remediation_rejected_before_staging(api, route, suffix, quality):
    client, settings = api
    before = set(settings.job_dir.rglob("*"))
    response = await client.post(route + quality, files={
        "file": (f"legacy.{suffix}", bytes.fromhex("d0cf11e0a1b11ae1")),
    })
    assert response.status_code == 415
    assert "docx" in response.json()["detail"].lower()
    assert set(settings.job_dir.rglob("*")) == before


@pytest.mark.parametrize("damage", ["missing", "empty", "directory", "unreadable"])
async def test_done_artifact_unavailable_returns_gone(api, tmp_path, monkeypatch, damage):
    client, settings = api
    assert (await client.get("/v1/jobs/unknown/result")).status_code == 404
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Retained sentinel"])
    job = await submit(client, "/v1/office/remediate", source.name, source.read_bytes())
    assert job["status"] == "done"
    output = settings.job_dir / job["id"] / "remediated.docx"
    if damage == "missing":
        output.unlink()
    elif damage == "empty":
        output.write_bytes(b"")
    elif damage == "directory":
        output.unlink()
        output.mkdir()
    else:
        from pathlib import Path
        original = Path.open
        def denied(path, *args, **kwargs):
            if path == output:
                raise PermissionError("private path must not escape")
            return original(path, *args, **kwargs)
        monkeypatch.setattr(Path, "open", denied)
    response = await client.get(f"/v1/jobs/{job['id']}/result")
    assert response.status_code == 410
    assert response.json() == {"detail": "Result artifact is no longer available."}


@pytest.mark.parametrize("fault", ["failure-with-file", "missing", "empty", "directory", "invalid"])
async def test_claimed_office_success_requires_readable_artifact(api, tmp_path, monkeypatch, fault):
    from project_remedy.office_remediator import OfficeRemediator, OfficeRemediationResult

    async def broken(self, input_path, output_path, **kwargs):
        if fault == "failure-with-file":
            make_docx(output_path, body_paragraphs=["Partial output"])
        elif fault == "empty":
            output_path.write_bytes(b"")
        elif fault == "directory":
            output_path.mkdir()
        elif fault == "invalid":
            output_path.write_bytes(b"PK\x03\x04not a package")
        return OfficeRemediationResult(input_path, output_path, success=fault != "failure-with-file")

    monkeypatch.setattr(OfficeRemediator, "remediate", broken)
    client, _ = api
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Sentinel"])
    job = await submit(client, "/v1/office/remediate", source.name, source.read_bytes())
    assert job["status"] == "failed"
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 409


async def test_advisory_findings_keep_genuine_output(api, tmp_path):
    client, _ = api
    # Generic hyperlink text is a real residual finding this remediator does not fix.
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Sentinel"],
                       hyperlinks=[("click here", "https://example.invalid/")])
    job = await submit(client, "/v1/office/remediate", source.name, source.read_bytes())
    assert job["status"] == "done"
    meta = json.loads(job["metadata_json"])
    assert meta["acceptance"]["package_valid"] is True
    assert meta["acceptance"]["passed"] is False
    assert meta["acceptance"]["failed_rule_ids"]
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 200


@pytest.mark.parametrize("advisory", ["acceptance", "quality"])
async def test_advisory_exception_keeps_output_without_disclosing_error(api, tmp_path, monkeypatch, advisory):
    def unavailable(*args, **kwargs):
        raise RuntimeError("private-source /private/path fixture-secret")
    if advisory == "acceptance":
        monkeypatch.setattr("project_remedy.office_acceptance.evaluate_office_acceptance", unavailable)
    else:
        monkeypatch.setattr("backend.app.quality_calibration.assert_quality_calibrated", unavailable)
    client, _ = api
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Sentinel"])
    job = await submit(client, "/v1/office/remediate?quality=true", source.name, source.read_bytes())
    assert job["status"] == "done"
    assert "fixture-secret" not in json.dumps(job)
    assert "unavailable" in job["metadata_json"].lower()
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 200


async def test_render_failure_then_success_isolated_in_worker(api, monkeypatch):
    from project_remedy.html_to_pdf import HTMLToPDFConverter, PDFConversionResult

    async def no_browser(self):
        pass
    async def failed(self, html, output_path, **kwargs):
        return PDFConversionResult(error_message="private renderer failure")
    monkeypatch.setattr(HTMLToPDFConverter, "start", no_browser)
    monkeypatch.setattr(HTMLToPDFConverter, "close", no_browser)
    monkeypatch.setattr(HTMLToPDFConverter, "convert", failed)
    client, _ = api
    job = await submit(client, "/v1/convert/html-to-pdf", "source.html", b"<html><body>Sentinel</body></html>")
    assert job["status"] == "failed"
    assert "private" not in job["error"]
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 409

    async def rendered(self, html, output_path, **kwargs):
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(str(output_path))
        canvas.drawString(72, 720, "Retained sentinel")
        canvas.save()
        return PDFConversionResult(output_path=output_path, success=True)
    monkeypatch.setattr(HTMLToPDFConverter, "convert", rendered)
    job = await submit(client, "/v1/convert/html-to-pdf", "valid.html", b"<html><body>Retained sentinel</body></html>")
    assert job["status"] == "done", job
    result = await client.get(f"/v1/jobs/{job['id']}/result")
    assert result.status_code == 200
    assert result.headers["content-type"] == "application/pdf"
    from pypdf import PdfReader
    assert "Retained sentinel" in PdfReader(BytesIO(result.content)).pages[0].extract_text()


@pytest.mark.parametrize("fault", ["missing", "empty", "invalid", "zero-pages", "failure-with-file", "exception"])
async def test_renderer_cannot_publish_unusable_pdf(api, monkeypatch, fault):
    import pikepdf
    from project_remedy.html_to_pdf import HTMLToPDFConverter, PDFConversionResult

    async def no_browser(self):
        pass
    async def broken(self, html, output_path, **kwargs):
        if fault == "exception":
            raise RuntimeError("private-source fixture-secret")
        if fault == "empty":
            output_path.write_bytes(b"")
        elif fault == "invalid":
            output_path.write_bytes(b"%PDF-1.7\nnot a PDF")
        elif fault in {"zero-pages", "failure-with-file"}:
            with pikepdf.new() as pdf:
                if fault == "failure-with-file":
                    pdf.add_blank_page()
                pdf.save(output_path)
        return PDFConversionResult(output_path=output_path, success=fault != "failure-with-file")
    monkeypatch.setattr(HTMLToPDFConverter, "start", no_browser)
    monkeypatch.setattr(HTMLToPDFConverter, "close", no_browser)
    monkeypatch.setattr(HTMLToPDFConverter, "convert", broken)
    client, _ = api
    job = await submit(client, "/v1/convert/html-to-pdf", "source.html", b"<html><body>Sentinel</body></html>")
    assert job["status"] == "failed"
    assert "fixture-secret" not in job["error"]
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 409


@pytest.mark.skipif(os.environ.get("REMEDY_FULL_RUNTIME_TESTS") != "1", reason="requires installed Chromium/full runtime")
async def test_real_html_render_download_after_failure(api):
    """No converter/renderer mocks: real Playwright + PDF post-processing."""
    client, _ = api
    # Admitted HTML with invalid UTF-8 fails in the real worker before rendering.
    invalid = await submit(client, "/v1/convert/html-to-pdf", "bad.html", b"<html><body>\xff</body></html>")
    assert invalid["status"] == "failed"
    assert (await client.get(f"/v1/jobs/{invalid['id']}/result")).status_code == 409
    job = await submit(client, "/v1/convert/html-to-pdf", "valid.html",
                       b"<!doctype html><html lang='en'><head><title>Fixture</title></head>"
                       b"<body><h1>Retained sentinel</h1></body></html>")
    assert job["status"] == "done", job
    result = await client.get(f"/v1/jobs/{job['id']}/result")
    assert result.status_code == 200
    from pypdf import PdfReader
    pdf = PdfReader(BytesIO(result.content))
    assert len(pdf.pages) >= 1
    assert "Retained sentinel" in pdf.pages[0].extract_text()


@pytest.fixture
def fake_ai(monkeypatch):
    """Intercept only the external HTTP transport, keeping extraction/generation real."""
    seen = []
    async def respond(self, request):
        assert request.url.host == "provider.invalid", request.url
        seen.append(request)
        prompt = json.loads(request.content)["messages"][-1]["content"]
        source = "Slide 1 Title" if "Slide 1 Title" in prompt else ("v2.1" if "v2.1" in prompt else "Retained sentinel")
        return httpx.Response(200, json={
            "choices": [{"message": {"content": f"<p>Retained sentinel {source}</p>"}}],
            "usage": {"prompt_tokens": 1, "completion_tokens": 1},
        }, request=request)
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", respond)
    return seen


@pytest.mark.parametrize("suffix", ["docx", "pptx", "xlsx"])
async def test_office_html_real_extraction_generation_after_corrupt_input(api, tmp_path, fake_ai, suffix):
    client, _ = api
    failed = await submit(client, "/v1/convert/office-to-html", f"bad.{suffix}", b"PK\x03\x04broken")
    assert failed["status"] == "failed"
    assert (await client.get(f"/v1/jobs/{failed['id']}/result")).status_code == 409
    source = office_fixture(tmp_path, suffix)
    job = await submit(client, "/v1/convert/office-to-html", source.name, source.read_bytes())
    assert job["status"] == "done", job
    result = await client.get(f"/v1/jobs/{job['id']}/result")
    assert result.status_code == 200
    assert result.headers["content-type"].startswith("text/html")
    from bs4 import BeautifulSoup
    html = BeautifulSoup(result.content, "html.parser")
    assert html.html and html.body
    assert "Retained sentinel" in html.body.get_text()
    # Prove native extraction reached the actual external planning call.
    plan = json.loads(fake_ai[0].content)["messages"][-1]["content"]
    sentinel = {"docx": "Retained sentinel", "pptx": "Slide 1 Title", "xlsx": "v2.1"}[suffix]
    assert sentinel in plan
    assert sentinel in html.body.get_text()


@pytest.mark.parametrize("claimed_html", ["", "not html", "<html><head></head></html>"])
async def test_office_html_invalid_claimed_output_fails(api, tmp_path, monkeypatch, fake_ai, claimed_html):
    from project_remedy.converter import HTMLConverter
    async def broken(self, job):
        return claimed_html
    monkeypatch.setattr(HTMLConverter, "convert", broken)
    client, _ = api
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Retained sentinel"])
    job = await submit(client, "/v1/convert/office-to-html", source.name, source.read_bytes())
    assert job["status"] == "failed"
    assert (await client.get(f"/v1/jobs/{job['id']}/result")).status_code == 409


async def test_queued_and_running_results_are_conflict(api, tmp_path, monkeypatch):
    from project_remedy.office_remediator import OfficeRemediator
    entered, release = asyncio.Event(), asyncio.Event()
    original = OfficeRemediator.remediate
    async def blocked(self, *args, **kwargs):
        entered.set()
        await release.wait()
        return await original(self, *args, **kwargs)
    monkeypatch.setattr(OfficeRemediator, "remediate", blocked)
    client, _ = api
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Sentinel"])
    try:
        running = await client.post("/v1/office/remediate", files={"file": (source.name, source.read_bytes())})
        await asyncio.wait_for(entered.wait(), 5)
        queued = await client.post("/v1/office/remediate", files={"file": (source.name, source.read_bytes())})
        for response, expected in [(running, "running"), (queued, "queued")]:
            assert response.status_code == 202
            job_id = response.json()["id"]
            assert (await client.get(f"/v1/jobs/{job_id}")).json()["status"] == expected
            assert (await client.get(f"/v1/jobs/{job_id}/result")).status_code == 409
    finally:
        release.set()
    assert (await terminal(client, running.json()["id"]))["status"] == "done"
    assert (await terminal(client, queued.json()["id"]))["status"] == "done"


async def test_generation_provider_error_then_valid_conversion(api, tmp_path, monkeypatch, fake_ai):
    transport = httpx.AsyncHTTPTransport.handle_async_request
    async def unavailable(self, request):
        raise httpx.ConnectError("fixture-secret private provider error", request=request)
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", unavailable)
    client, _ = api
    source = make_docx(tmp_path / "valid.docx", body_paragraphs=["Retained sentinel"])
    failed = await submit(client, "/v1/convert/office-to-html", source.name, source.read_bytes())
    assert failed["status"] == "failed"
    assert "fixture-secret" not in failed["error"]
    assert (await client.get(f"/v1/jobs/{failed['id']}/result")).status_code == 409
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", transport)
    valid = await submit(client, "/v1/convert/office-to-html", source.name, source.read_bytes())
    assert valid["status"] == "done"
    assert (await client.get(f"/v1/jobs/{valid['id']}/result")).status_code == 200


async def test_other_operations_keep_format_dispatch(api, fake_ai):
    client, _ = api
    ole = bytes.fromhex("d0cf11e0a1b11ae1")
    # Accepted legacy Office-to-HTML reaches its existing native extraction,
    # which fails. This compatibility check does not claim legacy support.
    job = await submit(client, "/v1/convert/office-to-html", "legacy.doc", ole)
    assert job["kind"] == "convert_office_to_html"
    assert job["status"] == "failed"
    for route in ["/v1/office/check", "/v1/convert/extract-markdown"]:
        response = await client.post(route, files={"file": ("legacy.doc", ole)})
        assert response.status_code != 415
    pdf = await client.post("/v1/remediate", files={"file": ("input.pdf", b"%PDF-1.7\nbroken")})
    assert pdf.status_code == 202
    assert pdf.json()["kind"] == "remediate_pdf"


async def test_result_contract_documented_and_auth_preserved(api):
    client, _ = api
    schema = (await client.get("/openapi.json")).json()
    responses = schema["paths"]["/v1/jobs/{job_id}/result"]["get"]["responses"]
    assert {"200", "404", "409", "410"} <= responses.keys()
    client.headers.pop("X-API-Key")
    assert (await client.get("/v1/jobs/unknown/result")).status_code == 401
